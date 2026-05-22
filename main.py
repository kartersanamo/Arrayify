import csv
import os
import re
from pathlib import Path
from docx import Document


REGISTER_NAME_RE = re.compile(r"^R(\d+)$")
TANK_DESCRIPTION_RE = re.compile(r"^(.*) @ (\d+)$")
REGISTER_ARRAY_NAME_RE = re.compile(r"^R(\d+)\[(\d+)\]$")
TABLE_TITLE_RE = re.compile(r"^Tank Soundings for (.+?) containing ")
FILES_FOLDER_NAME = "Files"
FILES_INPUT_FORMAT = "csv"
FILES_OUTPUT_FORMAT = "csv"
INPUT_FILE_NAME = "h221Test.csv"
INPUT_FILE_STEM = Path(INPUT_FILE_NAME).stem

TAG_PREFIX_OVERRIDES: dict[str, str | None] = {
    "Fuel Oil OverFlow Tank Volume": "FO_OVFL",
    "Fuel Oil SDT Day Tank Volume": "FO_SDT",
    "Ballast Anti Roll #1 Tank Volume": "BS_AR1",
    "Ballast Anti Roll #2 Tank Volume": "BS_AR2",
    "Ballast #7-C Tank Volume": "BS_7C",
    "Fuel Oil #6-P Tank Volume": "FO_6P",
    "Fuel Oil #6-S Tank Volume": "FO_6S",
    "Ballast AftPeak-P Tank Volume": "BS_APP",
    "Ballast AftPeak-S Tank Volume": "BS_APS",
    "Ballast ForePeak Tank Volume": "BS_FP",
    "Methanol Pump Void #1-S Volume": "ME_1S",
    "Methanol Pump Void #1-P Volume": "ME_1P",
    "Methanol Pump Void #2-S Volume": "ME_2S",
    "Methanol Pump Void #2-P Volume": "ME_2P",
    "WashWater Tank Volume": "WW",
    "Ballast AftPeak Void-P Tank Volume": "BS_APPV",
    "Ballast AftPeak Void-S Tank Volume": "BS_APSV",
}

def main() -> None:
    options: dict[int, list] = {
        1: [arrayify_points],
        2: [sound_tanks],
        3: [normalize_tags],
        4: [arrayify_points, sound_tanks, normalize_tags],
    }

    print("Program options:")
    print("- 1) Array-ify points")
    print("- 2) Sound tanks")
    print("- 3) Normalize tag names")
    print("- 4) All")
    option_choice: str = input("Enter your choice (1, 2, 3, 4): ")

    # Ensure they entered a number
    if not option_choice.isdigit():
        print("Please enter a number.")
        return
    
    option_choice_int: int = int(option_choice)

    # Ensured they entered a VALID number
    if not option_choice_int in options:
        print("Please enter a number 1-4.")
        return
    
    print("")

    for func in options.get(option_choice_int):
        func()

    return


def round_up_to_25(value: int) -> int:
    return ((value + 24) // 25) * 25


def is_register_name(value: str) -> bool:
    return bool(REGISTER_NAME_RE.match(value or ""))


def default_initial_value(data_type: str) -> str:
    if data_type.upper() == "REAL":
        return "0.0"
    return "0"


def load_tag_prefix_map() -> dict[str, str]:
    input_path = Path(FILES_FOLDER_NAME) / INPUT_FILE_NAME

    if not input_path.is_file():
        return {}

    with input_path.open(newline="") as input_csv:
        rows = list(csv.reader(input_csv, dialect="excel"))

    prefix_map: dict[str, str] = {}
    for row in rows[1:]:
        if len(row) <= 2:
            continue

        name = row[0].strip()
        description = row[2].strip()
        if not name or not description or not name.endswith("_LEVEL"):
            continue

        prefix_map.setdefault(canonicalize_tank_description(description), name[: -len("_LEVEL")])

    return prefix_map


def canonicalize_tank_description(description: str) -> str:
    canonical_description = description.strip()
    canonical_description = canonical_description.replace("Fuel Oil", "FO")
    canonical_description = canonical_description.replace("Liquid Mud", "LM")
    canonical_description = canonical_description.replace("Methanol", "METH")
    canonical_description = canonical_description.replace("Potable Water", "POTWATER")
    return canonical_description


def format_custom_tag_name(custom_tag: str) -> str | None:
    normalized_custom_tag = re.sub(r"[\s-]+", "_", custom_tag.strip())
    if not normalized_custom_tag:
        return None

    if normalized_custom_tag.endswith("_TANK_TABLE"):
        return normalized_custom_tag

    if normalized_custom_tag.endswith("_TABLE"):
        normalized_custom_tag = normalized_custom_tag[: -len("_TABLE")]

    if normalized_custom_tag.endswith("_TANK"):
        return f"{normalized_custom_tag}_TABLE"

    return f"{normalized_custom_tag}_TANK_TABLE"


def normalize_tag_name(
    row: list[str],
    prefix_map: dict[str, str],
    custom_tag_map: dict[str, str] | None = None,
) -> str | None:
    if len(row) <= 2:
        return None

    description = row[2].strip()
    if not description:
        return None

    description_match = TANK_DESCRIPTION_RE.match(description)
    base_description = description_match.group(1).strip() if description_match else description
    if custom_tag_map is not None and base_description in custom_tag_map:
        base_name = custom_tag_map[base_description]
        register_match = REGISTER_ARRAY_NAME_RE.match(row[0])
        if register_match:
            return f"{base_name}[{register_match.group(2)}]"

        if is_register_name(row[0]):
            return base_name

        return None

    override_prefix = TAG_PREFIX_OVERRIDES.get(base_description)
    if override_prefix is not None:
        prefix = override_prefix
    else:
        canonical_description = canonicalize_tank_description(base_description)
        if base_description in TAG_PREFIX_OVERRIDES and TAG_PREFIX_OVERRIDES[base_description] is None:
            return None
        prefix = prefix_map.get(canonical_description)
    if prefix is None:
        return None

    normalized_prefix = prefix if prefix.endswith("_TANK") else f"{prefix}_TANK"
    base_name = f"{normalized_prefix}_TABLE"
    register_match = REGISTER_ARRAY_NAME_RE.match(row[0])
    if register_match:
        return f"{base_name}[{register_match.group(2)}]"

    if is_register_name(row[0]):
        return base_name

    return None


def print_normalize_summary(
    matched_rows: list[dict[str, str]],
    unmatched_rows: list[dict[str, str]],
) -> None:
    print("Normalize summary:")

    def group_rows(rows: list[dict[str, str]]) -> list[dict[str, object]]:
        grouped_rows: dict[str, dict[str, object]] = {}
        for item in rows:
            key = item["description"].split(" @ ", 1)[0]
            group = grouped_rows.setdefault(
                key,
                {
                    "description": key,
                    "tag": item.get("tag", ""),
                    "rows": [],
                },
            )
            group["rows"].append(item)
        return list(grouped_rows.values())

    def print_grouped_rows(title: str, rows: list[dict[str, str]], include_tag: bool) -> None:
        if not rows:
            print(f"{title} none")
            return

        print(title)
        for group in group_rows(rows):
            sample_rows = group["rows"][:2]
            count = len(group["rows"])
            if include_tag:
                print(f"- {group['description']} -> {group['tag']} ({count} rows)")
            else:
                print(f"- {group['description']} ({count} rows)")

            for index, item in enumerate(sample_rows):
                end = "\n" if index == 0 else "...\n\n"
                if include_tag:
                    print(f"  - {item['register']} | {item['description']} -> {item['tag']}", end = end)
                else:
                    print(f"  - {item['register']} | {item['description']}", end = end)

    print_grouped_rows("Renamed rows:", matched_rows, include_tag=True)
    print_grouped_rows("Rows without a tag match:", unmatched_rows, include_tag=False)


def build_array_row(
    source_row: list[str],
    base_register: int,
    base_description: str,
    index: int,
    is_padded: bool,
    padded_initial_value: str | None = None
) -> list[str]:
    row = source_row.copy()
    row[0] = f"R{base_register}[{index}]"
    row[2] = f"{base_description} @ {index}"
    row[15] = f"%R{base_register:05d}"

    if is_padded:
        row[12] = padded_initial_value if padded_initial_value is not None else default_initial_value(row[1])

    return row


def build_base_row(source_row: list[str], base_description: str, target_length: int) -> list[str]:
    row = source_row.copy()
    row[0] = row[0].split("[", 1)[0]
    row[2] = base_description
    row[7] = str(target_length)
    row[12] = ", ".join(["0"] * target_length)
    return row


def extract_table_key(table_title: str) -> str | None:
    match = TABLE_TITLE_RE.match(table_title.strip())
    if not match:
        return None

    return match.group(1).strip()


def table_key_to_description(table_key: str) -> str | None:
    fuel_match = re.match(r"^FUEL(\d)\.(P|S|C)$", table_key)
    if fuel_match:
        number = fuel_match.group(1)
        side = fuel_match.group(2)
        return f"Fuel Oil #{number}-{side} Tank Volume"

    if table_key == "FO_DAY.S":
        return "Fuel Oil Day-S Tank Volume"
    elif table_key == "FO_DAY.P":
        return "Fuel Oil Day-P Tank Volume"
    
    # Is this correct?
    #elif table_key == "FO-EMERG.S":
    #    return "Fuel Oil Dirty Oil Tank Volume"
    # Is this correct?
    #elif table_key == "FO_SWG.S":
    #    return "Fuel Oil SDT Day Tank Volume"
    
    elif table_key == "FO_OVER.S":
        return "Fuel Oil OverFlow Tank Volume"

    lm_match = re.match(r"^LM(\d)\.(P|S)$", table_key)
    if lm_match:
        number = lm_match.group(1)
        side = lm_match.group(2)
        return f"Liquid Mud #{number}-{side} Tank Volume"

    meth_match = re.match(r"^METH(\d)\.(P|S)$", table_key)
    if meth_match:
        number = meth_match.group(1)
        side = meth_match.group(2)
        return f"Methanol #{number}-{side} Tank Volume"

    pot_match = re.match(r"^POTWATER\.(P|S)$", table_key)
    if pot_match:
        side = pot_match.group(1)
        return f"Potable Water-{side} Tank Volume"

    if table_key == "SEWAGE.S":
        return "Sewage Tank Volume"
    if table_key == "WASHWATER.P":
        return "WashWater Tank Volume"

    if table_key == "FOREPEAK.C":
        return "Ballast ForePeak Tank Volume"

    ballast_match = re.match(r"^BALLAST(\d+)([CW])\.(P|S|C)$", table_key)
    if ballast_match:
        number = ballast_match.group(1)
        location = ballast_match.group(2)
        side = ballast_match.group(3)
        if location == "W":
            return f"Ballast #{number}{side}-W Tank Volume"
        return f"Ballast #{number}{side}-C Tank Volume"

    if table_key == "ANTIROLL1.C":
        return "Ballast Anti Roll #1 Tank Volume"
    if table_key == "ANTIROLL2.C":
        return "Ballast Anti Roll #2 Tank Volume"
    if table_key == "BALLAST1.P":
        return None
    if table_key == "BALLAST1.S":
        return "Ballast #1-S Tank Volume"
    if table_key == "BALLAST7.C":
        return "Ballast #7-C Tank Volume"
    if table_key == "AFTPEAK1.P":
        return "Ballast AftPeak-P Tank Volume"
    if table_key == "AFTPEAK1.S":
        return "Ballast AftPeak-S Tank Volume"

    return None


def read_sounding_volumes(table) -> list[str]:
    volumes: list[str] = []
    for row in table.rows[2:]:
        if len(row.cells) < 2:
            continue

        volume_text = row.cells[1].text.strip()
        if not volume_text:
            continue

        volumes.append(volume_text)

    return volumes


def print_sound_mapping_report(
    matched_rows: list[dict[str, str]],
    unmatched_doc_tables: list[dict[str, str]],
    unmatched_csv_rows: list[str],
) -> None:
    print("Sound mapping report:")

    if matched_rows:
        print("Matched tanks:")
        for item in matched_rows:
            print(f"- {item['csv_register']} | {item['csv_description']} -> {item['doc_file']} :: {item['doc_title']}")
    else:
        print("Matched tanks: none")

    if unmatched_doc_tables:
        print("\nDOCX tables without a CSV match:")
        for item in unmatched_doc_tables:
            print(f"- {item['doc_file']} :: {item['doc_title']}")
    else:
        print("DOCX tables without a CSV match: none")

    if unmatched_csv_rows:
        print("\nCSV tank rows without a DOCX table:")
        for description in unmatched_csv_rows:
            print(f"- {description}")
    else:
        print("CSV tank rows without a DOCX table: none")


def print_arrayify_summary(summary_rows: list[dict[str, int | str]]) -> None:
    print("Arrayify summary:")
    print(f"Total tanks found: {len(summary_rows)}")

    total_points_found = 0
    total_points_allocated = 0

    for item in summary_rows:
        total_points_found += int(item["points_found"])
        total_points_allocated += int(item["points_allocated"])
        print(
            f"- {item['register']} | {item['description']} | "
            f"found {item['points_found']} -> allocated {item['points_allocated']} "
            f"(+{item['points_allocated'] - item['points_found']})"
        )

    print(f"Total points found: {total_points_found}")
    print(f"Total points allocated: {total_points_allocated}")
    print(f"Total padding added: {total_points_allocated - total_points_found}")


def set_block_initial_values(rows: list[list[str]], start_index: int, values: list[str]) -> None:
    base_row = rows[start_index]
    base_register = base_row[0].split("[", 1)[0]

    block_end = start_index + 1
    while block_end < len(rows):
        register_match = REGISTER_ARRAY_NAME_RE.match(rows[block_end][0])
        if not register_match or register_match.group(1) != base_register[1:]:
            break
        block_end += 1

    existing_length = block_end - start_index - 1
    if existing_length <= 0:
        raise ValueError(f"No array rows found for {base_row[2]}.")
    if not values:
        raise ValueError(f"No sounding values found for {base_row[2]}.")

    target_length = max(existing_length, round_up_to_25(len(values)))
    padded_values = values + [values[-1]] * (target_length - len(values))

    if target_length > existing_length:
        last_row_template = rows[block_end - 1]
        extra_rows: list[list[str]] = []

        for index in range(existing_length, target_length):
            row = last_row_template.copy()
            row[0] = f"{base_register}[{index}]"
            row[2] = f"{base_row[2]} @ {index}"
            row[12] = padded_values[index]
            row[15] = f"%R{int(base_register[1:]) + index:05d}"
            extra_rows.append(row)

        rows[block_end:block_end] = extra_rows
        block_end += len(extra_rows)

    base_row[7] = str(target_length)
    base_row[12] = ", ".join(padded_values)

    for offset, row in enumerate(rows[start_index + 1:block_end]):
        row[12] = padded_values[offset]


def find_base_row_index(rows: list[list[str]], description: str) -> int | None:
    for row_index, row in enumerate(rows[1:], start=1):
        if len(row) <= 12:
            continue

        if row[2] != description:
            continue

        if not is_register_name(row[0]) or '[' in row[0]:
            continue

        return row_index

    return None


def normalize_tags() -> None:
    print("Normalizing tags....")

    source_candidates = [
        Path(FILES_FOLDER_NAME) / f"{INPUT_FILE_STEM}_sounded.csv",
        Path(FILES_FOLDER_NAME) / f"{INPUT_FILE_STEM}_arrayified.csv",
    ]
    source_path = next((path for path in source_candidates if path.is_file()), None)
    if source_path is None:
        print("No arrayified or sounded CSV found. Run arrayify first.")
        return

    output_path = Path(FILES_FOLDER_NAME) / f"{INPUT_FILE_STEM}_normalized.csv"
    if output_path.exists():
        print(f"Conflicting output file path: {output_path}")
        return

    prefix_map = load_tag_prefix_map()
    if not prefix_map:
        print(f"Tag prefix map could not be built from {Path(FILES_FOLDER_NAME) / INPUT_FILE_NAME}.")
        return

    with source_path.open(newline="") as source_file:
        rows = list(csv.reader(source_file, dialect="excel"))

    if not rows:
        print(f"Source file is empty: {source_path}")
        return

    renamed_rows = 0
    matched_rows: list[dict[str, str]] = []
    unmatched_rows: list[dict[str, str]] = []
    custom_tag_map: dict[str, str] = {}
    prompted_descriptions: set[str] = set()

    for row in rows[1:]:
        if len(row) <= 2:
            continue

        source_register = row[0].split("[", 1)[0]
        description = row[2].strip()
        if not description:
            continue

        if is_register_name(row[0]) or REGISTER_ARRAY_NAME_RE.match(row[0]):
            normalized_name = normalize_tag_name(row, prefix_map, custom_tag_map)
            if normalized_name is None:
                if description_match := TANK_DESCRIPTION_RE.match(description):
                    base_description = description_match.group(1).strip()
                else:
                    base_description = description

                if base_description not in prompted_descriptions:
                    prompted_descriptions.add(base_description)
                    custom_input = input(
                        f"No tag match for '{base_description}'. Enter a custom tag prefix, or press Enter to leave it unchanged: "
                    ).strip()

                    if custom_input:
                        custom_tag_name = format_custom_tag_name(custom_input)
                        if custom_tag_name is None:
                            print(f"Skipping '{base_description}' because the custom tag was empty.")
                        else:
                            custom_tag_map[base_description] = custom_tag_name
                            normalized_name = normalize_tag_name(row, prefix_map, custom_tag_map)

                if normalized_name is None:
                    unmatched_rows.append(
                        {
                            "register": source_register,
                            "description": description,
                        }
                    )
                    continue

            matched_rows.append(
                {
                    "register": source_register,
                    "description": description,
                    "tag": normalized_name,
                }
            )

            row[0] = normalized_name
            if len(row) > 15:
                row[15] = ""
            renamed_rows += 1
            continue

        normalized_name = normalize_tag_name(row, prefix_map, custom_tag_map)
        if normalized_name is None or normalized_name == row[0]:
            continue

        row[0] = normalized_name
        if len(row) > 15:
            row[15] = ""
        renamed_rows += 1

    print_normalize_summary(matched_rows, unmatched_rows)

    with output_path.open("w", newline="") as output_file:
        writer = csv.writer(output_file, dialect="excel")
        writer.writerows(rows)

    print(f"Normalized {renamed_rows} row names.")
    print(f"Wrote normalized CSV to {output_path}")


def sound_tanks() -> None:
    print("Sounding tanks....")

    folder_input = input("Enter the folder path containing DOCX soundings [Files/NewSounds]: ").strip()
    sound_folder = Path(folder_input or Path(FILES_FOLDER_NAME) / "NewSounds")

    if not sound_folder.is_dir():
        print(f"Folder not found: {sound_folder}")
        return

    template_path = Path(FILES_FOLDER_NAME) / f"{INPUT_FILE_STEM}_arrayified.csv"
    if not template_path.is_file():
        print(f"Template file not found: {template_path}")
        return

    output_name = f"{INPUT_FILE_STEM}_sounded.csv"
    output_path = Path(FILES_FOLDER_NAME) / output_name
    if output_path.exists():
        print(f"Conflicting output file path: {output_path}")
        return

    with template_path.open(newline="") as template_file:
        rows = list(csv.reader(template_file, dialect="excel"))

    if not rows:
        print("Template file is empty. You must arrayify the input file first.")
        return

    base_rows: list[tuple[int, str]] = []
    base_register_by_description: dict[str, str] = {}
    for row_index, row in enumerate(rows[1:], start=1):
        if len(row) <= 12 or not row[0] or '[' in row[0] or not is_register_name(row[0]):
            continue

        base_rows.append((row_index, row[2]))
        base_register_by_description[row[2]] = row[0].split("[", 1)[0]

    explicit_matches: dict[str, dict[str, object]] = {}
    unmatched_doc_tables: list[dict[str, str]] = []

    for doc_path in sorted(sound_folder.glob("*.docx")) + sorted(sound_folder.glob("*.DOCX")):
        document = Document(doc_path)
        for table in document.tables:
            table_title = table.cell(0, 0).text.strip().replace("\n", " ")
            table_key = extract_table_key(table_title)
            if not table_key:
                continue

            volumes = read_sounding_volumes(table)
            description = table_key_to_description(table_key)

            if description is None:
                unmatched_doc_tables.append(
                    {
                        "doc_file": doc_path.name,
                        "doc_title": table_title,
                        "table_key": table_key,
                    }
                )
                continue

            if description in explicit_matches:
                if explicit_matches[description]["volumes"] != volumes:
                    raise ValueError(f"Duplicate sounding table for {description}.")
                continue

            explicit_matches[description] = {
                "doc_file": doc_path.name,
                "doc_title": table_title,
                "table_key": table_key,
                "volumes": volumes,
            }

    matched_metadata_by_description: dict[str, dict[str, object]] = {}

    for description, metadata in explicit_matches.items():
        row_index = find_base_row_index(rows, description)
        if row_index is None:
            raise ValueError(f"No CSV row found for sounding table {description}.")

        set_block_initial_values(rows, row_index, metadata["volumes"])
        matched_metadata_by_description[description] = metadata

    matched_rows: list[dict[str, str]] = []
    for _, description in base_rows:
        metadata = matched_metadata_by_description.get(description)
        if metadata is None:
            continue

        matched_rows.append(
            {
                "csv_register": base_register_by_description[description],
                "csv_description": description,
                "doc_file": str(metadata["doc_file"]),
                "doc_title": str(metadata["doc_title"]),
            }
        )

    used_doc_pairs = {
        (str(metadata["doc_file"]), str(metadata["doc_title"]))
        for metadata in matched_metadata_by_description.values()
    }
    unmatched_doc_tables = [
        item
        for item in unmatched_doc_tables
        if (item["doc_file"], item["doc_title"]) not in used_doc_pairs
    ]
    unmatched_csv_rows = [description for _, description in base_rows if description not in matched_metadata_by_description]

    print_sound_mapping_report(matched_rows, unmatched_doc_tables, unmatched_csv_rows)

    with output_path.open("w", newline="") as output_file:
        writer = csv.writer(output_file, dialect="excel")
        writer.writerows(rows)

    print(f"Wrote sounded CSV to {output_path}")

def arrayify_points() -> None:
    print("Array-ify-ing...")

    input_name = os.path.join("Files", INPUT_FILE_NAME)
    output_name = os.path.join("Files", f"{INPUT_FILE_STEM}_arrayified.csv")

    if not os.path.isfile(input_name):
        print(f"Input file not found: {input_name}")
        return

    if os.path.exists(output_name):
        print(f"Conflicting output file path: {output_name}")
        return

    with open(input_name, newline="") as input_csv:
        rows = list(csv.reader(input_csv, dialect="excel"))

    if not rows:
        print("Input file is empty.")
        return

    output_rows: list[list[str]] = [rows[0]]
    summary_rows: list[dict[str, int | str]] = []
    row_index = 1

    while row_index < len(rows):
        row = rows[row_index]
        if len(row) <= 15 or not is_register_name(row[0]):
            row_index += 1
            continue

        description_match = TANK_DESCRIPTION_RE.match(row[2])
        if not description_match or description_match.group(2) != "0":
            row_index += 1
            continue

        base_description = description_match.group(1)
        base_register = int(row[0][1:])

        block_rows: list[list[str]] = []
        expected_index = 0
        scan_index = row_index

        while scan_index < len(rows):
            current_row = rows[scan_index]
            if len(current_row) <= 15 or not is_register_name(current_row[0]):
                break

            current_description_match = TANK_DESCRIPTION_RE.match(current_row[2])
            if not current_description_match:
                break

            current_register = int(current_row[0][1:])
            current_description_index = int(current_description_match.group(2))

            if (
                current_description_match.group(1) != base_description
                or current_description_index != expected_index
                or current_register != base_register + expected_index
            ):
                break

            block_rows.append(current_row)
            expected_index += 1
            scan_index += 1

        if len(block_rows) <= 1:
            row_index += 1
            continue

        target_length = round_up_to_25(len(block_rows))
        last_real_initial_value = block_rows[-1][12]

        summary_rows.append(
            {
                "register": row[0],
                "description": base_description,
                "points_found": len(block_rows),
                "points_allocated": target_length,
            }
        )

        output_rows.append(
            build_base_row(
                source_row=block_rows[0],
                base_description=base_description,
                target_length=target_length,
            )
        )

        for index in range(target_length):
            if index < len(block_rows):
                source_row = block_rows[index]
                is_padded = False
            else:
                source_row = block_rows[-1]
                is_padded = True

            output_rows.append(
                build_array_row(
                    source_row=source_row,
                    base_register=base_register,
                    base_description=base_description,
                    index=index,
                    is_padded=is_padded,
                    padded_initial_value=last_real_initial_value,
                )
            )

        row_index = scan_index

    print_arrayify_summary(summary_rows)

    with open(output_name, "w", newline="") as output_csv:
        writer = csv.writer(output_csv, dialect="excel")
        writer.writerows(output_rows)

    print(f"Wrote {len(output_rows) - 1} modified rows to {output_name}")

if __name__ == "__main__":
    main()