from __future__ import annotations

from pathlib import Path

from docx import Document

from tank_tools.config import ProjectConfig
from tank_tools.io import CsvRepository
from tank_tools.models import (
    ArrayifySummaryRow,
    NormalizeMatch,
    NormalizeMiss,
    SoundDocumentMiss,
    SoundMatch,
)
from tank_tools.rules import TankRules


class ArrayifyService:
    def __init__(self, config: ProjectConfig, csv_repository: CsvRepository, rules: TankRules) -> None:
        self._config = config
        self._csv_repository = csv_repository
        self._rules = rules

    def arrayify_points(
        self,
        input_path: Path | None = None,
        output_path: Path | None = None,
        event_callback: callable | None = None,
    ) -> list[list[str]] | None:
        print("Array-ify-ing...")

        input_path = input_path or self._config.input_csv_path
        output_path = output_path or self._config.arrayified_csv_path

        if event_callback is not None:
            event_callback({"type": "status", "message": "Starting arrayify workflow."})

        if not input_path.is_file():
            print(f"Input file not found: {input_path}")
            return

        if output_path.exists():
            print(f"Conflicting output file path: {output_path}")
            return

        rows = self._csv_repository.read_rows(input_path)
        if not rows:
            print("Input file is empty.")
            return

        output_rows: list[list[str]] = [rows[0]]
        summary_rows: list[ArrayifySummaryRow] = []
        row_index = 1

        while row_index < len(rows):
            row = rows[row_index]
            if len(row) <= 15 or not self._rules.is_register_name(row[0]):
                row_index += 1
                continue

            description_match = self._rules.tank_description_re.match(row[2])
            if not description_match or description_match.group(2) != "0":
                row_index += 1
                continue

            base_description = description_match.group(1)
            base_register = int(row[0][1:])

            block_rows: list[list[str]] = []
            expected_index = 0
            scan_index = row_index

            while scan_index < len(rows):
                current_row = rows[scan_index]
                if len(current_row) <= 15 or not self._rules.is_register_name(current_row[0]):
                    break

                current_description_match = self._rules.tank_description_re.match(current_row[2])
                if not current_description_match:
                    break

                current_register = int(current_row[0][1:])
                current_description_index = int(current_description_match.group(2))

                if (
                    current_description_match.group(1) != base_description
                    or current_description_index != expected_index
                    or current_register != base_register + expected_index
                ):
                    break

                block_rows.append(current_row)
                expected_index += 1
                scan_index += 1

            if len(block_rows) <= 1:
                row_index += 1
                continue

            target_length = self._rules.round_up_to_25(len(block_rows))
            last_real_initial_value = block_rows[-1][12]

            summary_rows.append(
                ArrayifySummaryRow(
                    register=row[0],
                    description=base_description,
                    points_found=len(block_rows),
                    points_allocated=target_length,
                )
            )

            output_rows.append(self._build_base_row(block_rows[0], base_description, target_length))

            if event_callback is not None:
                event_callback(
                    {
                        "type": "preview",
                        "workflow": "arrayify",
                        "register": row[0],
                        "description": base_description,
                        "rows": output_rows.copy(),
                    }
                )

            for index in range(target_length):
                source_row = block_rows[index] if index < len(block_rows) else block_rows[-1]
                output_rows.append(
                    self._build_array_row(
                        source_row=source_row,
                        base_register=base_register,
                        base_description=base_description,
                        index=index,
                        is_padded=index >= len(block_rows),
                        padded_initial_value=last_real_initial_value,
                    )
                )

            row_index = scan_index

        self._print_summary(summary_rows)
        self._csv_repository.write_rows(output_path, output_rows)
        print(f"Wrote {len(output_rows) - 1} modified rows to {output_path}")

        if event_callback is not None:
            event_callback({"type": "completed", "workflow": "arrayify", "rows": output_rows.copy()})

        return output_rows

    def _build_array_row(
        self,
        source_row: list[str],
        base_register: int,
        base_description: str,
        index: int,
        is_padded: bool,
        padded_initial_value: str | None = None,
    ) -> list[str]:
        row = source_row.copy()
        row[0] = f"R{base_register}[{index}]"
        row[2] = f"{base_description} @ {index}"
        row[15] = f"%R{base_register:05d}"

        if is_padded:
            row[12] = padded_initial_value if padded_initial_value is not None else self._rules.default_initial_value(row[1])

        return row

    @staticmethod
    def _build_base_row(source_row: list[str], base_description: str, target_length: int) -> list[str]:
        row = source_row.copy()
        row[0] = row[0].split("[", 1)[0]
        row[2] = base_description
        row[7] = str(target_length)
        row[12] = ", ".join(["0"] * target_length)
        return row

    @staticmethod
    def _print_summary(summary_rows: list[ArrayifySummaryRow]) -> None:
        print("Arrayify summary:")
        print(f"Total tanks found: {len(summary_rows)}")

        total_points_found = 0
        total_points_allocated = 0

        for item in summary_rows:
            total_points_found += item.points_found
            total_points_allocated += item.points_allocated
            print(
                f"- {item.register} | {item.description} | "
                f"found {item.points_found} -> allocated {item.points_allocated} "
                f"(+{item.points_allocated - item.points_found})"
            )

        print(f"Total points found: {total_points_found}")
        print(f"Total points allocated: {total_points_allocated}")
        print(f"Total padding added: {total_points_allocated - total_points_found}")


class TankSoundingService:
    def __init__(self, config: ProjectConfig, csv_repository: CsvRepository, rules: TankRules) -> None:
        self._config = config
        self._csv_repository = csv_repository
        self._rules = rules

    def sound_tanks(self) -> None:
        print("Sounding tanks....")

        folder_input = input("Enter the folder path containing DOCX soundings [Files/NewSounds]: ").strip()
        sound_folder = Path(folder_input or self._config.new_sounds_dir)

        if not sound_folder.is_dir():
            print(f"Folder not found: {sound_folder}")
            return

        if not self._config.arrayified_csv_path.is_file():
            print(f"Template file not found: {self._config.arrayified_csv_path}")
            return

        if self._config.sounded_csv_path.exists():
            print(f"Conflicting output file path: {self._config.sounded_csv_path}")
            return

        rows = self._csv_repository.read_rows(self._config.arrayified_csv_path)
        if not rows:
            print("Template file is empty. You must arrayify the input file first.")
            return

        base_rows: list[tuple[int, str]] = []
        base_register_by_description: dict[str, str] = {}
        for row_index, row in enumerate(rows[1:], start=1):
            if len(row) <= 12 or not row[0] or "[" in row[0] or not self._rules.is_register_name(row[0]):
                continue

            base_rows.append((row_index, row[2]))
            base_register_by_description[row[2]] = row[0].split("[", 1)[0]

        explicit_matches: dict[str, dict[str, object]] = {}
        unmatched_doc_tables: list[SoundDocumentMiss] = []

        for doc_path in sorted(sound_folder.glob("*.docx")) + sorted(sound_folder.glob("*.DOCX")):
            document = Document(doc_path)
            for table in document.tables:
                table_title = table.cell(0, 0).text.strip().replace("\n", " ")
                table_key = self._rules.extract_table_key(table_title)
                if not table_key:
                    continue

                volumes = self._rules.read_sounding_volumes(table)
                description = self._rules.table_key_to_description(table_key)

                if description is None:
                    unmatched_doc_tables.append(SoundDocumentMiss(doc_file=doc_path.name, doc_title=table_title))
                    continue

                if description in explicit_matches:
                    if explicit_matches[description]["volumes"] != volumes:
                        raise ValueError(f"Duplicate sounding table for {description}.")
                    continue

                explicit_matches[description] = {
                    "doc_file": doc_path.name,
                    "doc_title": table_title,
                    "table_key": table_key,
                    "volumes": volumes,
                }

        matched_metadata_by_description: dict[str, dict[str, object]] = {}
        for description, metadata in explicit_matches.items():
            row_index = self._rules.find_base_row_index(rows, description)
            if row_index is None:
                raise ValueError(f"No CSV row found for sounding table {description}.")

            self._set_block_initial_values(rows, row_index, metadata["volumes"])
            matched_metadata_by_description[description] = metadata

        matched_rows: list[SoundMatch] = []
        for _, description in base_rows:
            metadata = matched_metadata_by_description.get(description)
            if metadata is None:
                continue

            matched_rows.append(
                SoundMatch(
                    csv_register=base_register_by_description[description],
                    csv_description=description,
                    doc_file=str(metadata["doc_file"]),
                    doc_title=str(metadata["doc_title"]),
                )
            )

        used_doc_pairs = {
            (str(metadata["doc_file"]), str(metadata["doc_title"]))
            for metadata in matched_metadata_by_description.values()
        }
        unmatched_doc_tables = [
            item for item in unmatched_doc_tables if (item.doc_file, item.doc_title) not in used_doc_pairs
        ]
        unmatched_csv_rows = [description for _, description in base_rows if description not in matched_metadata_by_description]

        self._print_sound_mapping_report(matched_rows, unmatched_doc_tables, unmatched_csv_rows)
        self._csv_repository.write_rows(self._config.sounded_csv_path, rows)
        print(f"Wrote sounded CSV to {self._config.sounded_csv_path}")

    def _set_block_initial_values(self, rows: list[list[str]], start_index: int, values: list[str]) -> None:
        base_row = rows[start_index]
        base_register = base_row[0].split("[", 1)[0]

        block_end = start_index + 1
        while block_end < len(rows):
            register_match = self._rules.register_array_name_re.match(rows[block_end][0])
            if not register_match or register_match.group(1) != base_register[1:]:
                break
            block_end += 1

        existing_length = block_end - start_index - 1
        if existing_length <= 0:
            raise ValueError(f"No array rows found for {base_row[2]}.")
        if not values:
            raise ValueError(f"No sounding values found for {base_row[2]}.")

        target_length = max(existing_length, self._rules.round_up_to_25(len(values)))
        padded_values = values + [values[-1]] * (target_length - len(values))

        if target_length > existing_length:
            last_row_template = rows[block_end - 1]
            extra_rows: list[list[str]] = []

            for index in range(existing_length, target_length):
                row = last_row_template.copy()
                row[0] = f"{base_register}[{index}]"
                row[2] = f"{base_row[2]} @ {index}"
                row[12] = padded_values[index]
                row[15] = f"%R{int(base_register[1:]) + index:05d}"
                extra_rows.append(row)

            rows[block_end:block_end] = extra_rows
            block_end += len(extra_rows)

        base_row[7] = str(target_length)
        base_row[12] = ", ".join(padded_values)

        for offset, row in enumerate(rows[start_index + 1:block_end]):
            row[12] = padded_values[offset]

    @staticmethod
    def _print_sound_mapping_report(
        matched_rows: list[SoundMatch],
        unmatched_doc_tables: list[SoundDocumentMiss],
        unmatched_csv_rows: list[str],
    ) -> None:
        print("Sound mapping report:")

        if matched_rows:
            print("Matched tanks:")
            for item in matched_rows:
                print(f"- {item.csv_register} | {item.csv_description} -> {item.doc_file} :: {item.doc_title}")
        else:
            print("Matched tanks: none")

        if unmatched_doc_tables:
            print("\nDOCX tables without a CSV match:")
            for item in unmatched_doc_tables:
                print(f"- {item.doc_file} :: {item.doc_title}")
        else:
            print("DOCX tables without a CSV match: none")

        if unmatched_csv_rows:
            print("\nCSV tank rows without a DOCX table:")
            for description in unmatched_csv_rows:
                print(f"- {description}")
        else:
            print("CSV tank rows without a DOCX table: none")


class TagNormalizationService:
    def __init__(self, config: ProjectConfig, csv_repository: CsvRepository, rules: TankRules) -> None:
        self._config = config
        self._csv_repository = csv_repository
        self._rules = rules

    def normalize_tags(self) -> None:
        print("Normalizing tags....")

        source_path = self._find_source_path()
        if source_path is None:
            print("No arrayified or sounded CSV found. Run arrayify first.")
            return

        if self._config.normalized_csv_path.exists():
            print(f"Conflicting output file path: {self._config.normalized_csv_path}")
            return

        rows = self._csv_repository.read_rows(source_path)
        if not rows:
            print(f"Source file is empty: {source_path}")
            return

        prefix_map = self._rules.load_tag_prefix_map(self._csv_repository.read_rows(self._config.input_csv_path))
        if not prefix_map:
            print(f"Tag prefix map could not be built from {self._config.input_csv_path}.")
            return

        renamed_rows = 0
        matched_rows: list[NormalizeMatch] = []
        unmatched_rows: list[NormalizeMiss] = []
        custom_tag_map: dict[str, str] = {}
        prompted_descriptions: set[str] = set()

        for row in rows[1:]:
            if len(row) <= 2:
                continue

            source_register = row[0].split("[", 1)[0]
            description = row[2].strip()
            if not description:
                continue

            if self._rules.is_register_name(row[0]) or self._rules.register_array_name_re.match(row[0]):
                normalized_name = self._rules.normalize_tag_name(row, prefix_map, custom_tag_map)
                if normalized_name is None:
                    base_description = self._rules.extract_base_description(description)
                    if base_description not in prompted_descriptions:
                        prompted_descriptions.add(base_description)
                        custom_input = input(
                            f"No tag match for '{base_description}'. Enter a custom tag prefix, or press Enter to leave it unchanged: "
                        ).strip()

                        if custom_input:
                            custom_tag_name = self._rules.format_custom_tag_name(custom_input)
                            if custom_tag_name is None:
                                print(f"Skipping '{base_description}' because the custom tag was empty.")
                            else:
                                custom_tag_map[base_description] = custom_tag_name
                                normalized_name = self._rules.normalize_tag_name(row, prefix_map, custom_tag_map)

                    if normalized_name is None:
                        unmatched_rows.append(NormalizeMiss(register=source_register, description=description))
                        continue

                matched_rows.append(NormalizeMatch(register=source_register, description=description, tag=normalized_name))
                row[0] = normalized_name
                if len(row) > 15:
                    row[15] = ""
                renamed_rows += 1
                continue

            normalized_name = self._rules.normalize_tag_name(row, prefix_map, custom_tag_map)
            if normalized_name is None or normalized_name == row[0]:
                continue

            row[0] = normalized_name
            if len(row) > 15:
                row[15] = ""
            renamed_rows += 1

        self._print_normalize_summary(matched_rows, unmatched_rows)
        self._csv_repository.write_rows(self._config.normalized_csv_path, rows)
        print(f"Normalized {renamed_rows} row names.")
        print(f"Wrote normalized CSV to {self._config.normalized_csv_path}")

    def _find_source_path(self) -> Path | None:
        candidates = [self._config.sounded_csv_path, self._config.arrayified_csv_path]
        return next((path for path in candidates if path.is_file()), None)

    @staticmethod
    def _print_normalize_summary(
        matched_rows: list[NormalizeMatch],
        unmatched_rows: list[NormalizeMiss],
    ) -> None:
        print("Normalize summary:")

        def group_rows(rows: list[NormalizeMatch | NormalizeMiss]) -> list[dict[str, object]]:
            grouped_rows: dict[str, dict[str, object]] = {}
            for item in rows:
                key = item.description.split(" @ ", 1)[0]
                group = grouped_rows.setdefault(
                    key,
                    {
                        "description": key,
                        "tag": getattr(item, "tag", ""),
                        "rows": [],
                    },
                )
                group["rows"].append(item)
            return list(grouped_rows.values())

        def print_grouped_rows(title: str, rows: list[NormalizeMatch | NormalizeMiss], include_tag: bool) -> None:
            if not rows:
                print(f"{title} none")
                return

            print(title)
            for group in group_rows(rows):
                sample_rows = group["rows"][:2]
                count = len(group["rows"])
                if include_tag:
                    print(f"- {group['description']} -> {group['tag']} ({count} rows)")
                else:
                    print(f"- {group['description']} ({count} rows)")

                for index, item in enumerate(sample_rows):
                    end = "\n" if index == 0 else "...\n\n"
                    if include_tag:
                        print(f"  - {item.register} | {item.description} -> {item.tag}", end=end)
                    else:
                        print(f"  - {item.register} | {item.description}", end=end)

        print_grouped_rows("Renamed rows:", matched_rows, include_tag=True)
        print_grouped_rows("Rows without a tag match:", unmatched_rows, include_tag=False)
